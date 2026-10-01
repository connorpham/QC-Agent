"""Synthetic fixture files generated at test time; nothing binary is committed and no customer
content is used."""

import io
import stat
import zipfile
from collections.abc import Awaitable, Callable
from pathlib import Path

import openpyxl
import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation


def make_docx(path: Path, headings: list[tuple[int, str]], paragraphs: list[str]) -> Path:
    document = DocxDocument()
    for level, text in headings:
        document.add_heading(text, level=level)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]]) -> Path:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for title, body in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.save(str(path))
    return path


def make_xlsx(path: Path, sheet: str, rows: list[list[str]]) -> Path:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    workbook.save(str(path))
    return path


def make_pdf(path: Path, pages: list[str], *, password: str | None = None) -> Path:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        y = 72
        for line in text.splitlines() or [""]:
            page.insert_text((72, y), line, fontsize=11)
            y += 14
    if password:
        document.save(
            str(path),
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            user_pw=password,
            owner_pw=password,
        )
    else:
        document.save(str(path))
    document.close()
    return path


def make_zip(
    path: Path,
    entries: dict[str, bytes],
    *,
    symlink: str | None = None,
) -> Path:
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
        if symlink is not None:
            info = zipfile.ZipInfo(symlink)
            info.create_system = 3
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
            archive.writestr(info, "target.txt")
    return path


def pdf_bytes(pages: list[str]) -> bytes:
    document = pymupdf.open()
    for text in pages:
        document.new_page().insert_text((72, 72), text, fontsize=11)
    data = document.tobytes()
    document.close()
    return bytes(data)


def docx_bytes(paragraphs: list[str]) -> bytes:
    buffer = io.BytesIO()
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def bytes_reader(data: bytes) -> Callable[[int], Awaitable[bytes]]:
    """Build an async chunked-reader over ``data``, mimicking ``UploadFile.read``: each call
    returns up to ``size`` bytes and an empty ``bytes`` once exhausted."""
    view = memoryview(data)
    position = 0

    async def read(size: int) -> bytes:
        nonlocal position
        chunk = bytes(view[position : position + size])
        position += size
        return chunk

    return read
